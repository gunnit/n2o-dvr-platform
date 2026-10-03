"""HACCP process order, flow diagram and manual output (segnalazione
2026-10-02 "da rivedere passaggi HACCP")."""

from __future__ import annotations

import asyncio
import importlib.util
import sys
from pathlib import Path

from docx import Document

from app.data.haccp_activity_types import (
    ccp_process_key,
    get_default_ccps,
    is_prerequisite,
    process_flow,
)

BACKEND_ROOT = Path(__file__).resolve().parents[1]


def test_catalog_ccps_sort_along_the_process_with_prerequisites_last():
    ordered = [c["nome"] for c in sorted(get_default_ccps("mensa_aziendale"), key=ccp_process_key)]
    assert ordered == [
        "Ricevimento materie prime",
        "Conservazione a freddo",
        "Scongelamento",
        "Cottura",
        "Abbattimento temperatura",
        "Trasporto pasti",
        "Igiene del personale",
        "Pulizia e sanificazione superfici",
    ]


def test_order_follows_the_name_not_the_code():
    """CCPs added by hand or by the AI number themselves freely."""
    ccps = [
        {"codice": "CCP1", "nome": "Cottura"},
        {"codice": "CCP9", "nome": "Ricevimento merci"},
        {"codice": "CUSTOM-1", "nome": "Abbattimento"},
    ]
    assert [c["nome"] for c in sorted(ccps, key=ccp_process_key)] == [
        "Ricevimento merci", "Cottura", "Abbattimento",
    ]


def test_igiene_and_sanificazione_are_prerequisites():
    assert is_prerequisite({"codice": "CCP7", "nome": "Igiene del personale", "fase": "Prima di ogni lavorazione"})
    assert is_prerequisite({"codice": "X", "nome": "Sanificazione attrezzature"})
    assert not is_prerequisite({"codice": "CCP1", "nome": "Cottura"})


def test_flow_always_has_receipt_storage_and_service():
    flow = process_flow(get_default_ccps("supermercato_retail"), "supermercato_retail")
    assert flow[0] == "Ricevimento merci"
    assert flow[-1] == "Somministrazione / vendita"
    assert "Preparazione e lavorazione" not in flow
    flow = process_flow(get_default_ccps("ristorante_con_cucina"), "ristorante_con_cucina")
    assert flow.index("Cottura") < flow.index("Abbattimento della temperatura")


def test_manual_prints_flow_full_ccp_table_and_no_donor_specifics(tmp_path):
    spec = importlib.util.spec_from_file_location(
        "verify_all_generators", str(BACKEND_ROOT / "scripts" / "verify_all_generators.py")
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    fixture = module.build_fixture()
    fixture["haccp_config"].tipologia_attivita = "mensa_aziendale"
    fixture["haccp_config"].ccps = get_default_ccps("mensa_aziendale")
    module.patch_generators(fixture, str(tmp_path))
    ok, path, message = asyncio.run(module.run_one("HACCP", fixture["azienda"].id))
    assert ok, message

    doc = Document(path)
    text = "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    )
    for donor in ("BARONI", "Prodotti di panetteria per la vendita", "BACONE BAR"):
        assert donor not in text, donor
    assert "Diagramma di flusso delle fasi" in text
    assert "Prerequisiti igienici" in text
    # Fields the operator enters are printed (they used to be dropped).
    assert "Respingere la fornitura" in text  # azione correttiva of CCP4
    assert "Ogni consegna" in text  # frequenza of CCP4
    # The SOP index matches the procedures the body documents.
    assert "SOP 21" in text and "SOP 25" not in text
    assert "Maria Conti" in text  # responsabile in place of the donor's name
