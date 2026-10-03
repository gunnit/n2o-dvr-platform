"""Allegato stress: one valutazione per mansione (segnalazione 2026-10-02).

N2O asked to assess every mansione. The allegato used to load a single
arbitrary stress row (``.limit(1)``, no order), so per-mansione assessments
never reached the document.
"""

from __future__ import annotations

import asyncio
import importlib.util
import sys
from pathlib import Path

from docx import Document

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))


def _load_verify():
    spec = importlib.util.spec_from_file_location(
        "verify_all_generators",
        str(BACKEND_ROOT / "scripts" / "verify_all_generators.py"),
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _text(path: str) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def _generate(module, fixture, tmp_path) -> str:
    module.patch_generators(fixture, str(tmp_path))
    ok, path, message = asyncio.run(
        module.run_one("ALLEGATO_STRESS", fixture["azienda"].id)
    )
    assert ok, message
    return path


def test_every_mansione_is_printed_with_a_summary(tmp_path):
    module = _load_verify()
    fixture = module.build_fixture()
    generale = fixture["stress"]
    cuoco = module.mk(
        gruppo_omogeneo="Azienda intera", mansione="Cuoco",
        area_a_eventi_sentinella={}, area_b_contenuto_lavoro={}, area_c_contesto_lavoro={},
        punteggio_a=2, punteggio_b=12, punteggio_c=15, punteggio_totale=29,
        livello_rischio="MEDIO", misure_correttive=None,
    )
    cameriere = module.mk(
        gruppo_omogeneo="Azienda intera", mansione="Cameriere",
        area_a_eventi_sentinella={}, area_b_contenuto_lavoro={}, area_c_contesto_lavoro={},
        punteggio_a=0, punteggio_b=4, punteggio_c=5, punteggio_totale=9,
        livello_rischio="BASSO", misure_correttive="Riunioni trimestrali.",
    )
    fixture["stress_all"] = [generale, cameriere, cuoco]

    text = _text(_generate(module, fixture, tmp_path))

    assert "Quadro riepilogativo per mansione" in text
    assert "Valutazione generale (azienda intera)" in text
    assert "Mansione: Cameriere" in text
    assert "Mansione: Cuoco" in text
    assert "29 / 67" in text and "9 / 67" in text
    assert "Riunioni trimestrali." in text


def test_single_valutazione_has_no_summary_table(tmp_path):
    module = _load_verify()
    fixture = module.build_fixture()
    text = _text(_generate(module, fixture, tmp_path))
    assert "Quadro riepilogativo per mansione" not in text
    assert "Valutazione generale (azienda intera)" in text
