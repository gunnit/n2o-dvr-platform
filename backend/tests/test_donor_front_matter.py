"""The donor allegati arrive filled in, not as blank forms.

``ALLEGATO RISCHIO INCENDIO.docx`` and ``PIANO GESTIONE EMERGENZE - AZIENDA.docx``
are N2O's blank master forms. Until 2026-09-11 the generators appended their
assessment after the template without ever writing into it, so the client
opened 28 pages of empty boxes — anagrafica, roster, safety roles, six
per-ambiente sheets — asking for data the platform already holds.

These tests pin both halves of the fix: what must now be filled, and what
must no longer be printed at all. They also fail if a template is replaced
with one whose labels differ, which is the only way the fills can silently
stop working.
"""

import asyncio
import importlib.util
import sys
from pathlib import Path

import pytest
from docx import Document

from app.services.document_generator.allegato_incendio import SUPERSEDED_CHAPTERS

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))


def _load_verify():
    spec = importlib.util.spec_from_file_location(
        "verify_all_generators",
        str(BACKEND_ROOT / "scripts" / "verify_all_generators.py"),
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _generate(tmp_path, name: str) -> Document:
    module = _load_verify()
    fixture = module.build_fixture()
    module.patch_generators(fixture, str(tmp_path))
    ok, path, msg = asyncio.run(module.run_one(name, fixture["azienda"].id))
    assert ok, msg
    return Document(path)


def _text(doc: Document) -> str:
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def _blank_forms(doc: Document) -> list[str]:
    """Two-column tables whose value column is empty in every row."""
    blank = []
    for table in doc.tables:
        rows = list(table.rows)
        if len(rows) < 3 or len(table.columns) != 2:
            continue
        values = [
            (row.cells[1].text or "").strip()
            for row in rows
            if len(row.cells) > 1 and row.cells[1]._tc is not row.cells[0]._tc
        ]
        if values and not any(values):
            blank.append((rows[0].cells[0].text or "").strip()[:40])
    return blank


@pytest.fixture(scope="module")
def incendio(tmp_path_factory) -> Document:
    return _generate(tmp_path_factory.mktemp("incendio"), "ALLEGATO_INCENDIO")


@pytest.fixture(scope="module")
def pee(tmp_path_factory) -> Document:
    return _generate(tmp_path_factory.mktemp("pee"), "PEE_AZIENDA")


@pytest.mark.parametrize("name", ["incendio", "pee"])
def test_no_blank_donor_forms(name, request):
    doc = request.getfixturevalue(name)
    assert _blank_forms(doc) == []


@pytest.mark.parametrize("name", ["incendio", "pee"])
def test_anagrafica_is_filled_from_the_survey(name, request):
    doc = request.getfixturevalue(name)
    anagrafica = next(
        t for t in doc.tables
        if len(t.columns) == 2
        and {(r.cells[0].text or "").strip() for r in t.rows} >= {"Azienda", "Datore di Lavoro"}
    )
    values = {
        (row.cells[0].text or "").strip(): (row.cells[1].text or "").strip()
        for row in anagrafica.rows
    }
    assert values["Azienda"] == "ACME MECCANICA COMPOSITA SRL"
    assert values["Datore di Lavoro"] == "Mario Rossi"
    assert values["Responsabile del Servizio di Prevenzione e Protezione (RSPP)"] == "Luca Bianchi"
    assert values["Rappresentanti dei Lavoratori per la Sicurezza"] == "Giulia Verdi"
    # Nobody in the roster is a medico competente: a fact about the survey,
    # not a gap the reader has to guess at.
    assert values["Medico Competente"] == "Non nominato"
    # The platform models neither, so the operator is asked — not invented.
    assert values["Addetto del Servizio di Prevenzione e Protezione (ASPP)"] == "[DA COMPILARE]"
    assert values["Dirigente per la sicurezza"] == "[DA COMPILARE]"


@pytest.mark.parametrize("name", ["incendio", "pee"])
def test_roster_and_roles_are_filled(name, request):
    text = _text(request.getfixturevalue(name))
    for nominativo in ("Mario Rossi", "Luca Bianchi", "Giulia Verdi", "Antonio Marrone"):
        assert nominativo in text
    # The addetto al primo soccorso reaches his own donor sheet.
    assert "Operaio Tornitore" in text


def test_ambienti_index_lists_the_survey(incendio):
    index = next(
        t for t in incendio.tables
        if t.rows and [(c.text or "").strip() for c in t.rows[0].cells[:2]] == ["n.", "Ambiente di Lavoro"]
    )
    listed = [(row.cells[1].text or "").strip() for row in list(index.rows)[1:]]
    assert "Officina meccanica" in listed
    assert "" not in listed, "the ruled spare lines must be trimmed, not printed empty"


def test_superseded_donor_chapters_are_gone(incendio):
    text = _text(incendio)
    for heading in SUPERSEDED_CHAPTERS:
        assert heading not in text, f"{heading!r} is re-emitted from the live valutazione"
    # Including the cached index lines that pointed at them.
    toc = [p.text for p in incendio.paragraphs if (p.style.name or "").lower().startswith("toc")]
    assert not [line for line in toc if line.startswith(SUPERSEDED_CHAPTERS[0])]
    # The assessment the generator does emit is still there.
    assert "Valutazione per ambiente" in text
    assert "Misure di prevenzione e protezione per area" in text
    # ...and the removal stopped at the chapters it was aimed at. The index
    # lines carry page numbers, so they never match a chapter anchor; this
    # fails loudly if one ever does and the front matter goes with it.
    for kept in ("Introduzione", "Anagrafica Aziendale", "Criteri adottati", "Dichiarazione"):
        assert kept in text


def test_declaration_is_signed_with_the_company(incendio):
    text = _text(incendio)
    assert "Il sottoscritto, Mario Rossi in qualità di Datore di Lavoro della ACME MECCANICA COMPOSITA SRL" in text
    assert "___________ in qualità di" not in text


def test_fields_refresh_on_open(incendio):
    """Word must rebuild the donor's cached index, whose page numbers and
    chapter list both predate this generation."""
    settings = incendio.settings.element
    tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}updateFields"
    assert settings.find(tag) is not None
