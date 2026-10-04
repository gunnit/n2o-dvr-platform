"""DUVRI template body filled from the client's data (segnalazione
2026-10-02 "rivedere la generazione output del documento DUVRI")."""

from __future__ import annotations

import asyncio
import importlib.util
import re
import sys
from pathlib import Path

from docx import Document

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))


def _generate(tmp_path, mutate=None) -> Document:
    spec = importlib.util.spec_from_file_location(
        "verify_all_generators", str(BACKEND_ROOT / "scripts" / "verify_all_generators.py")
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    fixture = module.build_fixture()
    if mutate:
        mutate(fixture)
    module.patch_generators(fixture, str(tmp_path))
    ok, path, message = asyncio.run(module.run_one("DUVRI", fixture["azienda"].id))
    assert ok, message
    return Document(path)


def _all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def _table(doc: Document, first_cell: str):
    return next(t for t in doc.tables if t.rows and t.rows[0].cells[0].text.strip().startswith(first_cell))


def test_donor_identity_and_local_numbers_are_gone(tmp_path):
    text = _all_text(_generate(tmp_path))
    for donor in ("PESCHIERA", "Peschiera", "San Donato", "02 553 8300", "Farmacia Comunale", "RECOM"):
        assert donor not in text, donor
    assert "Numero Unico di Emergenza" in text


def test_anagrafica_attivita_and_dates_are_filled(tmp_path):
    doc = _generate(tmp_path)
    committente = next(
        t for t in doc.tables
        if t.rows and t.rows[0].cells[0].text.strip() == "Ragione Sociale"
    )
    values = {r.cells[0].text.strip(): r.cells[1].text.strip() for r in committente.rows}
    assert values["Ragione Sociale"] == "ACME MECCANICA COMPOSITA SRL"
    assert values["Datore di Lavoro"] == "Mario Rossi"

    text = _all_text(doc)
    assert "DATA FINE: DA DEFINIRE" not in text
    assert "DATA INIZIO: 01/05/2026" in text
    assert "(365 giorni)" in text
    assert "Il sottoscritto, Mario Rossi in qualità di datore di lavoro della ACME" in text
    assert not re.search(r"…{2,}", "\n".join(p.text for p in doc.paragraphs))


def test_workers_and_emergency_roles_come_from_the_censimento(tmp_path):
    doc = _generate(tmp_path)
    lavoratori = _table(doc, "Azienda Committente")
    nomi = [r.cells[1].text for r in lavoratori.rows[1:]]
    assert "Antonio Marrone" in nomi
    emergenze = next(
        t for t in doc.tables
        if t.rows and t.rows[0].cells[0].text.strip() == "Azienda"
        and t.rows[0].cells[1].text.strip().startswith("Ruolo")
    )
    assert emergenze.rows[1].cells[0].text == "ACME MECCANICA COMPOSITA SRL"
    assert emergenze.rows[3].cells[0].text == "Pulizie Industriali Parma SRL"


def test_fire_level_follows_the_incendio_assessment(tmp_path):
    def only_basso(fixture):
        for row in fixture["incendio"]:
            row.livello_rischio = "BASSO"

    doc = _generate(tmp_path, only_basso)
    cell = next(
        t for t in doc.tables
        if t.rows and t.rows[0].cells[0].text.startswith("Classificazione del livello di rischio incendio")
    ).rows[0].cells[1]
    assert cell.text.strip() == "BASSO"

    doc = _generate(tmp_path / "alto")
    cell = next(
        t for t in doc.tables
        if t.rows and t.rows[0].cells[0].text.startswith("Classificazione del livello di rischio incendio")
    ).rows[0].cells[1]
    assert cell.text.strip() == "ALTO"
