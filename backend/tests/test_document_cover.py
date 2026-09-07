"""The shared cover and running furniture every generated document carries.

Cover redesign 2026-09-07: letterhead strip, navy title band, identity and
revision tables, stamp/signature boxes; running header with the consultancy
mark and a two-line footer with the page number. Documents opened from a
donor template drop the donor's own cover first.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn

from app.services.document_generator.branding import Branding
from app.services.document_generator.design import (
    add_cover,
    add_running_header_footer,
    setup_document,
    strip_donor_cover,
)

BACKEND_ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = BACKEND_ROOT.parent / "templates"
GENERATED_AT = datetime(2026, 9, 7, 10, 0)


def _azienda(**overrides):
    fields = dict(
        ragione_sociale="Meccanica Brianza S.r.l.",
        sede_legale_via="Via dell'Industria 12",
        sede_legale_citta="Concorezzo",
        cap_legale="20863",
        provincia_legale="MB",
        sede_operativa_via="Via Manzoni 8",
        sede_operativa_citta="Vimercate",
        cap_operativa="20871",
        provincia_operativa="MB",
        partita_iva="01234567890",
        codice_ateco="25.62.00",
    )
    fields.update(overrides)
    return SimpleNamespace(**fields)


def _branding(**overrides) -> Branding:
    fields = dict(
        firm_name="N2O SRL",
        indirizzo="Via dei Chiosi 4",
        cap="20064",
        citta="Gorgonzola",
        provincia="MI",
        partita_iva="03726250131",
        telefono="02.91321649",
        email="sicurezza@n2o-sicurezza.it",
    )
    fields.update(overrides)
    return Branding(**fields)


def _text(doc: Document) -> str:
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def _blocks(doc: Document) -> list:
    return [el for el in doc.element.body if el.tag != qn("w:sectPr")]


def _has_picture(element) -> bool:
    return element.find(f".//{qn('w:drawing')}") is not None


def _new_doc() -> Document:
    doc = Document()
    setup_document(doc)
    return doc


def test_cover_carries_letterhead_band_identity_revision_and_signatures():
    doc = _new_doc()
    add_cover(
        doc,
        title="Allegato Rischio MMC",
        subtitle="Movimentazione Manuale dei Carichi",
        legal_basis="ai sensi del Titolo VI D.Lgs. 81/2008",
        azienda=_azienda(),
        branding=_branding(),
        version=3,
        generated_at=GENERATED_AT,
    )
    text = _text(doc)
    # Letterhead strip: the consultancy mark and its letterhead lines.
    assert _has_picture(doc.tables[0]._tbl)
    assert "N2O SRL" in text
    assert "Via dei Chiosi 4, 20064 Gorgonzola (MI)" in text
    assert "P.IVA 03726250131" in text
    # Title band, shaded navy.
    band = doc.tables[1]
    assert band.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "003D74"
    assert "ALLEGATO AL DOCUMENTO DI VALUTAZIONE DEI RISCHI" in text
    assert "Allegato Rischio MMC" in text
    assert "Movimentazione Manuale dei Carichi" in text
    # Company identity: name plus seats, VAT number and ATECO.
    assert "MECCANICA BRIANZA S.R.L." in text
    identity = doc.tables[2]
    assert [row.cells[0].text for row in identity.rows] == ["Sede legale", "Sede operativa", "Partita IVA", "Codice ATECO"]
    assert identity.rows[1].cells[1].text == "Via Manzoni 8, 20871 Vimercate (MB)"
    # Revision numbered like the Storico (version 3 -> Rev. 02, Aggiornamento).
    revisions = doc.tables[3]
    assert [cell.text for cell in revisions.rows[0].cells] == ["Rev.", "Motivazione", "Data"]
    assert [cell.text for cell in revisions.rows[1].cells] == ["02", "Aggiornamento", "07/09/2026"]
    # Stamp and signature boxes, then the page break that ends the cover.
    assert "TIMBRO E FIRMA DEL DATORE DI LAVORO" in text
    assert "FIRMA DELL'RSPP" in text
    assert _blocks(doc)[-1].find(f".//{qn('w:br')}").get(qn("w:type")) == "page"
    # The old "elaborato da" block is gone: the letterhead strip replaces it.
    assert "Documento elaborato da" not in text


def test_cover_without_letterhead_keeps_the_consultancy_off_the_page():
    doc = _new_doc()
    add_cover(
        doc,
        title="DOCUMENTO DI VALUTAZIONE DEI RISCHI",
        azienda=_azienda(),
        branding=_branding(firm_name="CONSULTANCY SENTINEL"),
        version=1,
        generated_at=GENERATED_AT,
        show_letterhead=False,
    )
    text = _text(doc)
    assert "CONSULTANCY SENTINEL" not in text
    assert "Via dei Chiosi" not in text
    assert not any(_has_picture(table._tbl) for table in doc.tables)


def test_cover_hero_image_and_fallback_marker(tmp_path):
    hero = BACKEND_ROOT / "assets" / "n2o_vera_dvr.png"
    doc = _new_doc()
    add_cover(
        doc,
        title="DVR",
        azienda=_azienda(),
        branding=_branding(),
        version=1,
        generated_at=GENERATED_AT,
        hero_image=str(hero),
        show_letterhead=False,
    )
    assert any(_has_picture(p._p) for p in doc.paragraphs)

    doc = _new_doc()
    add_cover(
        doc,
        title="DVR",
        azienda=_azienda(),
        branding=_branding(),
        version=1,
        generated_at=GENERATED_AT,
        hero_image=str(tmp_path / "missing.png"),
        hero_fallback_text="[LOGO N2O VERA NON DISPONIBILE]",
        show_letterhead=False,
    )
    assert "[LOGO N2O VERA NON DISPONIBILE]" in _text(doc)


def test_cover_degrades_to_nothing_on_sparse_data():
    doc = _new_doc()
    sparse = SimpleNamespace(ragione_sociale=None)
    add_cover(
        doc,
        title="Allegato",
        azienda=sparse,
        branding=Branding.default(),
        version=None,
        generated_at=GENERATED_AT,
    )
    text = _text(doc)
    assert "—" in text  # the company name slot, never a placeholder token
    assert "Sede legale" not in text and "Partita IVA" not in text
    assert "00" in text and "Emissione" in text


def test_cover_at_top_goes_in_front_of_existing_content_without_page_break():
    doc = _new_doc()
    doc.add_paragraph("Donor body paragraph")
    doc.add_table(rows=1, cols=1).rows[0].cells[0].text = "Donor table"
    added = add_cover(
        doc,
        title="Piano Operativo di Sicurezza",
        azienda=_azienda(),
        branding=_branding(),
        version=1,
        generated_at=GENERATED_AT,
        at_top=True,
        page_break=False,
    )
    blocks = _blocks(doc)
    assert added > 0
    assert blocks[0].tag == qn("w:tbl")  # the letterhead strip leads
    donor_paragraph = next(el for el in blocks if "".join(t.text or "" for t in el.iter(qn("w:t"))) == "Donor body paragraph")
    assert blocks.index(donor_paragraph) == added
    assert blocks[-1].find(f".//{qn('w:t')}").text == "Donor table"
    assert not any(el.find(f".//{qn('w:br')}") is not None for el in blocks[:added])


def test_strip_donor_cover_stops_at_the_section_break():
    doc = Document(str(TEMPLATES / "ALLEGATO RISCHIO INCENDIO.docx"))
    before = len(_blocks(doc))
    removed = strip_donor_cover(doc)
    blocks = _blocks(doc)
    assert removed == 11
    assert len(blocks) == before - removed
    assert blocks[0].find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
    assert "Documento di valutazione" not in _text(doc)[:200]
    assert doc.sections[0]._sectPr.find(qn("w:pgBorders")) is None


def test_strip_donor_cover_stops_at_a_marker_or_a_table():
    doc = Document(str(TEMPLATES / "HACCP.docx"))
    removed = strip_donor_cover(doc, stop_text="Definizioni")
    assert removed == 17
    assert doc.paragraphs[0].text == "Definizioni"

    doc = Document(str(TEMPLATES / "DUVRI.docx"))
    removed = strip_donor_cover(doc, stop_before_table=True)
    assert removed == 1
    assert _blocks(doc)[0].tag == qn("w:tbl")
    assert "DOCUMENTO UNICO" not in _text(doc)


def test_strip_donor_cover_clears_a_cover_only_body_and_leaves_unknown_templates_alone():
    doc = Document(str(TEMPLATES / "POS.docx"))
    assert strip_donor_cover(doc) == 0  # no boundary: nothing removed
    assert len(_blocks(doc)) == 10
    assert strip_donor_cover(doc, whole_body=True) == 10
    assert _blocks(doc) == []


def test_running_header_and_footer_show_mark_title_client_page_and_revision():
    doc = _new_doc()
    doc.add_paragraph("Body")
    add_running_header_footer(
        doc,
        title="Allegato Rischio MMC",
        azienda=_azienda(),
        branding=_branding(),
        version=1,
        generated_at=GENERATED_AT,
    )
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    assert _has_picture(header._p)
    assert "Allegato Rischio MMC" in header.text
    assert "Meccanica Brianza S.r.l." in header.text
    footer_lines = [p.text for p in section.footer.paragraphs]
    assert footer_lines[0].startswith("N2O SRL · Via dei Chiosi 4, 20064 Gorgonzola (MI) · P.IVA 03726250131")
    assert footer_lines[0].endswith("Pagina 1 di 1")
    assert footer_lines[1] == "Tel. 02.91321649 · sicurezza@n2o-sicurezza.it\tRev. 00 del 07/09/2026"
    instructions = [el.text.strip() for el in section.footer._element.iter(qn("w:instrText"))]
    assert instructions == ["PAGE", "NUMPAGES"]
    # The cover page stays clean.
    assert section.different_first_page_header_footer is True
    assert section.first_page_header.paragraphs[0].text == ""


def test_running_header_without_logo_and_margins_for_the_furniture():
    doc = Document(str(TEMPLATES / "PIANO GESTIONE EMERGENZE - AZIENDA.docx"))
    add_running_header_footer(
        doc,
        title="Piano di gestione delle emergenze",
        azienda=_azienda(),
        branding=Branding.default(),
        version=1,
        generated_at=GENERATED_AT,
        header_logo=False,
    )
    for section in doc.sections:
        assert not _has_picture(section.header._element)
        assert section.top_margin.cm >= 2.3
        assert section.bottom_margin.cm >= 2.4
    footer_lines = [p.text for p in doc.sections[0].footer.paragraphs]
    assert footer_lines[0] == "N2O SRL\tPagina 1 di 1"
    assert footer_lines[1] == "\tRev. 00 del 07/09/2026"
