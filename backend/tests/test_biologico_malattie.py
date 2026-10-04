"""Rischio biologico: DPI utilizzati and malattie contraibili with sintomi
and cura (segnalazione 2026-10-02)."""

from __future__ import annotations

import asyncio
import importlib.util
import sys
from pathlib import Path

import pytest
from docx import Document

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from app.services.document_generator.reference_data_biologico import (  # noqa: E402
    get_sector_defaults,
)


def _load_verify():
    spec = importlib.util.spec_from_file_location(
        "verify_all_generators",
        str(BACKEND_ROOT / "scripts" / "verify_all_generators.py"),
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _text(path: str) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


@pytest.mark.parametrize("settore", ["alimentare", "asilo", "dentisti"])
def test_every_default_agente_has_sintomi_and_cura(settore):
    defaults = get_sector_defaults(settore)
    assert defaults["dpi"]
    for a in defaults["agenti"]:
        assert a["sintomi"].strip() and a["cura"].strip(), a["nome"]


def test_unknown_settore_is_rejected():
    with pytest.raises(ValueError):
        get_sector_defaults("veterinari")


def _generate(module, fixture, tmp_path) -> str:
    module.patch_generators(fixture, str(tmp_path))
    ok, path, message = asyncio.run(
        module.run_one("allegato_biologico_alimentare", fixture["azienda"].id)
    )
    assert ok, message
    return path


def test_allegato_prints_malattie_and_dpi_entered_by_the_operator(tmp_path):
    module = _load_verify()
    fixture = module.build_fixture()
    row = fixture["biologico"][0]
    row.agenti_identificati = [
        {"nome": "Norovirus", "gruppo": "2", "via": "Oro-fecale",
         "patologia": "Gastroenterite acuta virale",
         "sintomi": "Vomito improvviso", "cura": "Reidratazione orale"},
    ]
    row.dpi_richiesti = [{"descrizione": "Guanti in nitrile"}]

    text = _text(_generate(module, fixture, tmp_path))

    assert "Malattie contraibili: sintomi e cura" in text
    assert "Vomito improvviso" in text and "Reidratazione orale" in text
    assert "DPI) utilizzati" in text
    assert "Guanti in nitrile" in text


def test_older_valutazioni_without_sintomi_fall_back_to_the_defaults(tmp_path):
    module = _load_verify()
    fixture = module.build_fixture()  # Salmonella saved without sintomi/cura
    text = _text(_generate(module, fixture, tmp_path))
    salmonella = next(
        a for a in get_sector_defaults("alimentare")["agenti"] if a["nome"] == "Salmonella spp."
    )
    assert salmonella["sintomi"] in text
