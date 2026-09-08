"""The allegato incendio prints the measures the operator reviewed on screen.

The checklist in the UI is served by ``/calculate/fire-measures`` from
``app/data/fire_measures.py``. A ``NULL`` selection means the operator left
the checklist at its default (every measure of the band ticked), a saved
selection is printed verbatim and an explicitly empty one says so. Before
this the generator carried a private list, so unticking a measure in the UI
never changed the document (UI/UX audit 2026-09-07, F5).
"""

import asyncio
import importlib.util
import sys
from pathlib import Path

from docx import Document

from app.data.fire_measures import get_measures_for_level
from app.services.document_generator.allegato_incendio import prescrizioni_per_area

BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

CUSTOM = "Verifica semestrale porte REI del magazzino"


def _load_verify():
    spec = importlib.util.spec_from_file_location(
        "verify_all_generators",
        str(BACKEND_ROOT / "scripts" / "verify_all_generators.py"),
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _document_text(path: str) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def _generate_with_selection(tmp_path, misure_prevenzione):
    module = _load_verify()
    fixture = module.build_fixture()
    # Keep one ALTO area so every canonical measure is attributable to it.
    row = next(r for r in fixture["incendio"] if r.livello_rischio == "ALTO")
    row.misure_prevenzione = misure_prevenzione
    fixture["incendio"] = [row]
    module.patch_generators(fixture, str(tmp_path))
    ok, path, msg = asyncio.run(
        module.run_one("ALLEGATO_INCENDIO", fixture["azienda"].id)
    )
    assert ok, msg
    return _document_text(path)


def test_prescrizioni_helper_expands_null_to_the_band_and_splits_selections():
    assert prescrizioni_per_area(None, "MEDIO") == get_measures_for_level("Medio")
    assert prescrizioni_per_area(None, None) == get_measures_for_level("Basso")
    assert prescrizioni_per_area(f" a \n\n b \n", "ALTO") == ["a", "b"]
    assert prescrizioni_per_area("", "ALTO") == []


def test_untouched_checklist_prints_every_canonical_measure_of_the_band(tmp_path):
    text = _generate_with_selection(tmp_path, None)
    for measure in get_measures_for_level("Alto"):
        assert measure in text
    # The generator's former private wording must not resurface.
    assert "livello 3-FOR" not in text


def test_saved_selection_replaces_the_canonical_list(tmp_path):
    canonical = get_measures_for_level("Alto")
    kept, dropped = canonical[1], canonical[0]
    text = _generate_with_selection(tmp_path, "\n".join([kept, CUSTOM]))
    assert kept in text
    assert CUSTOM in text
    assert dropped not in text
    assert "Misure aggiuntive registrate" not in text


def test_empty_selection_states_that_nothing_was_recorded(tmp_path):
    text = _generate_with_selection(tmp_path, "")
    assert "Nessuna prescrizione aggiuntiva registrata per quest'area." in text
    for measure in get_measures_for_level("Alto"):
        assert measure not in text
